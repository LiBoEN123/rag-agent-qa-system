"""小型真实 DOCX/HTML fixture，避免提交不必要的二进制或网页快照。"""

from io import BytesIO

from docx import Document

from app.ingestion.document_blocks import ParsedDocument, TableBlock, TableCell


REGISTRATION_HTML = """<html><body>
<p>办理说明</p>
<table><caption>企业登记办理要求</caption>
<thead><tr><th>事项</th><th>条件</th><th>材料</th></tr></thead>
<tbody><tr><td>企业注册</td><td>年满18岁</td><td>身份证、申请表</td></tr></tbody>
</table><p>其他说明</p>
</body></html>"""


ADMISSION_ROWS = (
    ("内蒙古工业大学2026年分省录取分数统计",),
    ("内蒙古本科提前B段艺术类",),
    ("专业名称", "最高分", "最低分", "平均分", "备注"),
    ("服装与服饰设计", "501", "465", "470", "投档成绩"),
    ("环境设计", "487", "475", "477.87", "投档成绩"),
    ("内蒙古本科批物理类",),
    ("专业名称", "最高分", "最低分", "平均分", "备注"),
    ("安全工程", "527", "501", "508.32", "物理类"),
    ("材料成型及控制工程", "549", "509", "517.98", "物理类"),
    ("材料成型及控制工程", "503", "480", "487.75", "民族班"),
    ("储能科学与工程", "566", "555", "556.69", "新能源学院"),
    ("储能科学与工程", "537", "529", "533.4", "兴边固边（新能源学院）"),
    ("储能科学与工程", "561", "544", "546.38", "物理类"),
)
ADMISSION_HTML = "<html><body><table>" + "".join(
    '<tr><td colspan="5">' + values[0] + "</td></tr>" if len(values) == 1 else
    "<tr>" + "".join(f"<{tag}>{value}</{tag}>" for value in values) + "</tr>"
    for row, values in enumerate(ADMISSION_ROWS)
    for tag in ["th" if row in (2, 6) else "td"]
) + "</table></body></html>"


def admission_structure() -> ParsedDocument:
    """与真实 Docling HTML 输出逐 cell 比对的分区表，纯 chunking 测试无需 parser。"""
    return ParsedDocument(blocks=(TableBlock(
        block_id="b0", order=0, row_count=len(ADMISSION_ROWS), column_count=5,
        cells=tuple(
            TableCell(text=text, row_index=row, column_index=column,
                      column_span=5 if len(values) == 1 else 1, is_header=row in (2, 6))
            for row, values in enumerate(ADMISSION_ROWS)
            for column, text in enumerate(values)
        ),
    ),))


def imut_admission_structure() -> ParsedDocument:
    """冻结 1741.htm 实抓 Docling b9 前 11 行；保留总计行及缺失的 header flags。

    2026-09-16 经 ExternalKnowledgeSourceService 抓取，原表 610×5。
    仅将 block_id/order 归零；row indices、文本、spans 和 flags 与原输出一致。
    """
    rows = (
        ("内蒙古工业大学2026年分省录取分数统计",),
        ("内蒙古本科提前B段艺术类",),
        ("专业名称", "最高分", "最低分", "平均分", "备注"),
        ("服装与服饰设计", "501", "465", "470", "投档成绩"),
        ("环境设计", "487", "475", "477.87", "投档成绩"),
        ("总计", "501", "465", "475.4", "投档成绩"),
        ("内蒙古本科批物理类",),
        ("专业名称", "最高分", "最低分", "平均分", "备注"),
        ("安全工程", "527", "501", "508.32", "物理类"),
        ("材料成型及控制工程", "549", "509", "517.98", "物理类"),
        ("材料成型及控制工程", "503", "480", "487.75", "民族班"),
    )
    return ParsedDocument(blocks=(TableBlock(
        block_id="b0", order=0, row_count=len(rows), column_count=5,
        cells=tuple(
            TableCell(text=text, row_index=row, column_index=column,
                      column_span=5 if len(values) == 1 else 1, is_header=False)
            for row, values in enumerate(rows)
            for column, text in enumerate(values)
        ),
    ),))


def registration_structure() -> ParsedDocument:
    """已持久化合同形态的中文表格，供 chunking 与真实索引/检索测试共用。"""
    values = [
        ("事项", "条件", "材料"),
        ("企业注册", "年满18岁", "身份证、申请表"),
        ("企业变更", "企业已注册", "营业执照、变更表"),
    ]
    return ParsedDocument(page_count=2, blocks=(TableBlock(
        block_id="b0", order=0, page_number=2, caption="企业登记办理要求",
        row_count=3, column_count=3,
        cells=tuple(
            TableCell(text=text, row_index=row, column_index=column, is_header=row == 0)
            for row, values_in_row in enumerate(values)
            for column, text in enumerate(values_in_row)
        ),
    ),))


def build_table_docx(*, merged: bool = False) -> bytes:
    document = Document()
    document.add_paragraph("段落 A")
    table = document.add_table(rows=3 if merged else 2, cols=3)
    table.style = "Table Grid"
    values = [("事项", "条件", "材料"), ("注册", "年满18岁", "身份证")]
    if merged:
        values.append(("续期", "成人", "证明"))
    for row, texts in zip(table.rows, values):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    if merged:
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(1, 2).merge(table.cell(2, 2))
    document.add_paragraph("段落 B")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
