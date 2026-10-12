"""Phase 12 protected ingestion operations API contracts。"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pydantic import SecretStr

from app.indexing.service import IndexingError
from app.api.application import create_app, get_default_ingestion_dependencies
import app.api.application as api_application
from app.api.ingestion import (
    INGESTION_ADMIN_HEADER,
    IngestionOperationsDependencies,
)
from app.ingestion import (
    IngestionCandidate,
    IngestionCandidateStatus,
    IngestionPrincipalType,
    IngestionRequester,
    IngestionService,
    build_knowledge_document,
)
from app.sources import (
    FirecrawlPagePayload,
    FirecrawlPageSourceAdapter,
    GitHubFilePayload,
    GitHubFileSourceAdapter,
    SourceIngestionError,
)
from app.storage.ingestion import (
    IndexRunOperation,
    IndexRunStatus,
    SQLiteCandidateRepository,
    SQLiteIndexRunRepository,
)
from app.storage.ingestion import _IngestionIndexRunRecord


ADMIN_KEY = "phase12-test-admin-key"
FIXED_TIME = datetime(2026, 9, 3, 11, 0, tzinfo=timezone.utc)
REQUESTER = IngestionRequester(
    principal_id="ingestion-admin-api",
    principal_type=IngestionPrincipalType.SERVICE,
    source="api",
)


class FakeSourceService:
    def __init__(self, ingestion: IngestionService) -> None:
        self.ingestion = ingestion
        self.proposal_contexts = []
        self.proposal_error = None

    async def propose_github_file(
        self,
        *,
        owner,
        repository,
        path,
        context,
        ref=None,
        reason=None,
    ):
        if self.proposal_error is not None:
            raise self.proposal_error
        self.proposal_contexts.append(context)
        document = GitHubFileSourceAdapter().to_document(
            GitHubFilePayload(
                owner=owner,
                repository=repository,
                path=path,
                requested_ref=ref,
                resolved_sha="a" * 40,
                content=("github body " + path) * 220,
                mime_type="text/markdown",
            ),
            fetched_at=FIXED_TIME,
        )
        return self.ingestion.propose(
            document=document,
            requester=REQUESTER,
            reason=reason,
            requested_at=FIXED_TIME,
        )

    async def propose_firecrawl_page(self, *, url, context, reason=None):
        if self.proposal_error is not None:
            raise self.proposal_error
        self.proposal_contexts.append(context)
        document = FirecrawlPageSourceAdapter().to_document(
            FirecrawlPagePayload(
                requested_url=url,
                final_url=url,
                title="Firecrawl page",
                content="firecrawl body",
                mime_type="text/html",
                status_code=200,
            ),
            fetched_at=FIXED_TIME,
        )
        return self.ingestion.propose(
            document=document,
            requester=REQUESTER,
            reason=reason,
            requested_at=FIXED_TIME,
        )



class FakeIndexingService:
    def __init__(self):
        self.index_calls = 0
        self.fail_index = False
        self.remove_calls = 0
        self.remove_error = None

    def remove_indexed(self, candidate_id):
        self.remove_calls += 1
        if self.remove_error:
            raise self.remove_error
        return {"num_added": 0, "num_updated": 0, "num_deleted": 1, "num_skipped": 0}

    def index_approved(self, candidate_id):
        self.index_calls += 1
        if self.fail_index:
            raise IndexingError(
                "SOURCE_EMBEDDINGS_UNAVAILABLE",
                "must not reach the API or database",
            )
        if self.index_calls == 1:
            return {
                "num_added": 1,
                "num_updated": 0,
                "num_deleted": 0,
                "num_skipped": 0,
            }
        return {
            "num_added": 0,
            "num_updated": 0,
            "num_deleted": 0,
            "num_skipped": 1,
        }


def _client(settings_factory, *, enabled=True):
    settings = settings_factory(
        ingestion_admin_api_enabled=enabled,
        ingestion_admin_api_key=ADMIN_KEY if enabled else None,
    )
    if not enabled:
        return TestClient(create_app(settings=settings)), None, None
    candidates = SQLiteCandidateRepository(settings.sqlite_db_path)
    ingestion = IngestionService(candidates)
    source = FakeSourceService(ingestion)
    runs = SQLiteIndexRunRepository(settings.sqlite_db_path)
    dependencies = IngestionOperationsDependencies(
        ingestion_service=ingestion,
        source_service=source,
        indexing_service=FakeIndexingService(),
        index_runs=runs,
        requester=REQUESTER,
        admin_api_key=SecretStr(ADMIN_KEY),
    )
    application = create_app(
        settings=settings,
        ingestion_dependencies=dependencies,
    )
    return TestClient(application), dependencies, runs


def _headers(key=ADMIN_KEY):
    return {INGESTION_ADMIN_HEADER: key}


def _github(client, path="README.md"):
    return client.post(
        "/admin/ingestion/sources/github",
        headers=_headers(),
        json={
            "owner": "owner",
            "repository": "repo",
            "path": path,
            "reason": "review",
        },
    )


def test_admin_routes_default_disabled_and_api_key_required(settings_factory, caplog):
    disabled, _, _ = _client(settings_factory, enabled=False)
    assert disabled.get("/admin/ingestion/candidates").status_code == 404

    client, _, _ = _client(settings_factory)
    missing = client.get("/admin/ingestion/candidates")
    wrong = client.get(
        "/admin/ingestion/candidates",
        headers=_headers("wrong-key"),
    )
    allowed = client.get("/admin/ingestion/candidates", headers=_headers())

    assert missing.status_code == 401
    assert wrong.status_code == 401
    assert missing.json()["error"]["code"] == "INGESTION_ADMIN_UNAUTHORIZED"
    assert allowed.status_code == 200
    assert ADMIN_KEY not in caplog.text
    assert ADMIN_KEY not in missing.text + wrong.text + allowed.text


def test_enabled_router_is_registered_during_production_lifespan(
    settings_factory,
    monkeypatch,
):
    settings = settings_factory(
        ingestion_admin_api_enabled=True,
        ingestion_admin_api_key=ADMIN_KEY,
    )
    monkeypatch.setattr(api_application, "get_settings", lambda: settings)
    monkeypatch.setattr(
        api_application,
        "start_application_resources",
        lambda _settings: None,
    )

    async def initialize_mcp(_settings):
        return None

    monkeypatch.setattr(api_application, "initialize_mcp_runtime", initialize_mcp)
    application = create_app()
    assert not any(
        route.path.startswith("/admin/ingestion")
        for route in application.routes
    )

    with TestClient(application) as client:
        response = client.get(
            "/admin/ingestion/candidates",
            headers=_headers(),
        )

    assert response.status_code == 200


def test_proposals_persist_safe_summaries_and_use_source_minimum_permission(
    settings_factory,
):
    client, source, _ = _client(settings_factory)
    github = _github(client)
    firecrawl = client.post(
        "/admin/ingestion/sources/firecrawl",
        headers=_headers(),
        json={"url": "https://example.com/article", "reason": "review"},
    )

    assert github.status_code == 201
    assert firecrawl.status_code == 201
    for response in (github, firecrawl):
        body = response.json()
        assert body["status"] == "pending"
        assert body["approval_status"] == "pending"
        assert "content" not in body
        assert "metadata" not in body
    assert source.source_service.proposal_contexts[0].permissions == ("mcp:github",)
    assert source.source_service.proposal_contexts[1].permissions == ("mcp:firecrawl",)
    assert all(item.principal_id == "ingestion-admin-api" for item in source.source_service.proposal_contexts)

    forbidden = client.post(
        "/admin/ingestion/sources/github",
        headers=_headers(),
        json={
            "owner": "owner",
            "repository": "repo",
            "path": "README.md",
            "content": "client controlled",
            "trust_level": "trusted_internal",
            "permissions": ["mcp:github", "mcp:firecrawl"],
        },
    )
    assert forbidden.status_code == 422


def test_stable_not_found_and_source_unavailable_error_mapping(settings_factory):
    client, source, _ = _client(settings_factory)
    missing = client.get(
        "/admin/ingestion/candidates/cand_missing",
        headers=_headers(),
    )
    source.source_service.proposal_error = SourceIngestionError(
        "SOURCE_TOOL_UNAVAILABLE",
        "internal provider detail",
    )
    unavailable = _github(client)

    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "CANDIDATE_NOT_FOUND"
    assert unavailable.status_code == 503
    assert unavailable.json()["error"]["code"] == "SOURCE_TOOL_UNAVAILABLE"
    assert "internal provider detail" not in unavailable.text


def test_candidate_list_filters_paginates_and_never_returns_content(settings_factory):
    client, _, _ = _client(settings_factory)
    first = _github(client, "a.md").json()
    second = _github(client, "b.md").json()
    client.post(
        f"/admin/ingestion/candidates/{first['candidate_id']}/approve",
        headers=_headers(),
        json={"reason": "approved"},
    )

    pending = client.get(
        "/admin/ingestion/candidates?status=pending&source_type=github&limit=1&offset=0",
        headers=_headers(),
    )
    approved = client.get(
        "/admin/ingestion/candidates?status=approved",
        headers=_headers(),
    )

    assert pending.status_code == 200
    assert [item["candidate_id"] for item in pending.json()] == [
        second["candidate_id"]
    ]
    assert approved.json()[0]["candidate_id"] == first["candidate_id"]
    assert "github body" not in pending.text + approved.text
    assert client.get(
        "/admin/ingestion/candidates?limit=101",
        headers=_headers(),
    ).status_code == 422


def test_candidate_detail_has_bounded_preview_and_latest_index_result(settings_factory):
    client, _, _ = _client(settings_factory)
    candidate = _github(client).json()
    detail = client.get(
        f"/admin/ingestion/candidates/{candidate['candidate_id']}",
        headers=_headers(),
    )

    assert detail.status_code == 200
    body = detail.json()
    assert len(body["content_preview"]) == 2000
    assert body["content_length"] > 2000
    assert body["preview_truncated"] is True
    assert body["content_is_untrusted"] is True
    assert body["latest_index_result"] is None
    assert client.get(
        f"/admin/ingestion/candidates/{candidate['candidate_id']}/full-content",
        headers=_headers(),
    ).status_code == 404


def test_approve_and_reject_are_explicit_terminal_operations(settings_factory):
    client, source, _ = _client(settings_factory)
    approved_candidate = _github(client, "approve.md").json()
    rejected_candidate = _github(client, "reject.md").json()

    approved = client.post(
        f"/admin/ingestion/candidates/{approved_candidate['candidate_id']}/approve",
        headers=_headers(),
        json={"reason": "verified"},
    )
    second_approve = client.post(
        f"/admin/ingestion/candidates/{approved_candidate['candidate_id']}/approve",
        headers=_headers(),
        json={"reason": "again"},
    )
    missing_reason = client.post(
        f"/admin/ingestion/candidates/{rejected_candidate['candidate_id']}/reject",
        headers=_headers(),
        json={},
    )
    rejected = client.post(
        f"/admin/ingestion/candidates/{rejected_candidate['candidate_id']}/reject",
        headers=_headers(),
        json={"reason": "not suitable"},
    )

    assert approved.json()["status"] == "approved"
    assert second_approve.status_code == 409
    assert missing_reason.status_code == 422
    assert rejected.json()["status"] == "rejected"
    assert source.indexing_service.index_calls == 0


def test_index_blocks_non_approved_persists_success_and_reindex(settings_factory):
    client, source, runs = _client(settings_factory)
    pending = _github(client, "pending.md").json()
    rejected = _github(client, "rejected.md").json()
    client.post(
        f"/admin/ingestion/candidates/{rejected['candidate_id']}/reject",
        headers=_headers(),
        json={"reason": "reject"},
    )

    for candidate in (pending, rejected):
        response = client.post(
            f"/admin/ingestion/candidates/{candidate['candidate_id']}/index",
            headers=_headers(),
        )
        assert response.status_code == 409
    assert runs.list() == []

    client.post(
        f"/admin/ingestion/candidates/{pending['candidate_id']}/approve",
        headers=_headers(),
        json={"reason": "approve"},
    )
    first = client.post(
        f"/admin/ingestion/candidates/{pending['candidate_id']}/index",
        headers=_headers(),
    )
    second = client.post(
        f"/admin/ingestion/candidates/{pending['candidate_id']}/index",
        headers=_headers(),
    )

    assert first.json()["status"] == "success"
    assert first.json()["added"] == 1
    assert second.json()["status"] == "success"
    assert second.json()["skipped"] == 1
    assert source.indexing_service.index_calls == 2
    listed = client.get(
        f"/admin/ingestion/index-runs?candidate_id={pending['candidate_id']}&status=success",
        headers=_headers(),
    )
    assert len(listed.json()) == 2


def test_index_failure_is_safe_and_candidate_remains_approved(settings_factory):
    client, source, runs = _client(settings_factory)
    candidate = _github(client, "failure.md").json()
    client.post(
        f"/admin/ingestion/candidates/{candidate['candidate_id']}/approve",
        headers=_headers(),
        json={"reason": "approve"},
    )
    source.indexing_service.fail_index = True

    response = client.post(
        f"/admin/ingestion/candidates/{candidate['candidate_id']}/index",
        headers=_headers(),
    )
    detail = client.get(
        f"/admin/ingestion/candidates/{candidate['candidate_id']}",
        headers=_headers(),
    )
    stored = runs.latest(candidate["candidate_id"])

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "SOURCE_EMBEDDINGS_UNAVAILABLE"
    assert "must not reach" not in response.text
    assert detail.json()["summary"]["status"] == "approved"
    assert detail.json()["latest_index_result"]["status"] == "failed"
    assert stored.status is IndexRunStatus.FAILED
    assert stored.error_message == "SOURCE_EMBEDDINGS_UNAVAILABLE"


def test_remove_requires_admin_and_retains_candidate_history(settings_factory):
    client, source, runs = _client(settings_factory)
    candidate = _github(client).json()
    path = f"/admin/ingestion/candidates/{candidate['candidate_id']}"
    for headers in ({}, _headers("wrong")):
        assert client.post(f"{path}/remove", headers=headers).status_code == 401
    assert client.post("/admin/ingestion/candidates/missing/remove", headers=_headers()).status_code == 404
    assert source.indexing_service.remove_calls == 0
    assert runs.list() == []
    before = source.ingestion_service.repository.get(candidate["candidate_id"])
    response = client.post(f"{path}/remove", headers=_headers())
    assert response.status_code == 200
    assert response.json()["operation"] == "remove"
    assert response.json()["status"] == "success"
    assert response.json()["deleted"] == 1
    assert source.ingestion_service.repository.get(candidate["candidate_id"]) == before
    detail = client.get(path, headers=_headers()).json()
    assert detail["latest_index_result"] == response.json()
    assert client.get("/admin/ingestion/index-runs?operation=remove", headers=_headers()).json() == [response.json()]
    assert client.get("/admin/ingestion/index-runs?operation=index", headers=_headers()).json() == []
    assert client.get("/admin/ingestion/index-runs?operation=delete", headers=_headers()).status_code == 422
    disabled, _, _ = _client(settings_factory, enabled=False)
    assert disabled.post(f"{path}/remove", headers=_headers()).status_code == 404


def _candidate_status(source, candidate_id, status):
    domain = source.ingestion_service
    if status == "approved":
        return domain.approve(candidate_id, decided_by=REQUESTER)
    if status == "rejected":
        return domain.reject(candidate_id, decided_by=REQUESTER, reason="review")
    candidate = domain.repository.get(candidate_id)
    if status == "superseded":
        values = candidate.model_dump()
        values["status"] = status
        return domain.repository.replace(IngestionCandidate.model_validate(values))
    return candidate


@pytest.mark.parametrize("status", list(IngestionCandidateStatus))
@pytest.mark.parametrize("removed", [False, True])
def test_delete_gate_uses_source_state_for_every_candidate_status(settings_factory, status, removed):
    client, source, runs = _client(settings_factory)
    summary = _github(client).json()
    candidate = _candidate_status(source, summary["candidate_id"], status)
    path = f"/admin/ingestion/candidates/{candidate.candidate_id}"
    if removed:
        assert client.post(f"{path}/remove", headers=_headers()).status_code == 200
    history = runs.list()
    assert client.get(path, headers=_headers()).json()["can_delete"] is removed
    response = client.delete(path, headers=_headers())
    if removed:
        assert response.status_code == 204
        assert response.content == b""
        assert source.ingestion_service.repository.get(candidate.candidate_id) is None
        assert client.get(path, headers=_headers()).status_code == 404
        assert client.get("/admin/ingestion/candidates", headers=_headers()).json() == []
    else:
        assert response.status_code == 409
        assert response.json()["error"]["code"] == "CANDIDATE_SOURCE_NOT_REMOVED"
        assert source.ingestion_service.repository.get(candidate.candidate_id) == candidate
    assert runs.list() == history
    assert source.indexing_service.remove_calls == int(removed)
    assert source.indexing_service.index_calls == 0


def test_delete_requires_auth_and_missing_candidate_is_404(settings_factory):
    client, source, runs = _client(settings_factory)
    path = "/admin/ingestion/candidates/missing"
    for headers in ({}, _headers("wrong")):
        assert client.delete(path, headers=headers).status_code == 401
    response = client.delete(path, headers=_headers())
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "CANDIDATE_NOT_FOUND"
    assert runs.list() == []
    assert source.indexing_service.index_calls == source.indexing_service.remove_calls == 0
    disabled, _, _ = _client(settings_factory, enabled=False)
    assert disabled.delete(path, headers=_headers()).status_code == 404


@pytest.mark.parametrize("operations,allowed", [
    ([("index", "success")], False),
    ([("remove", "failed")], False),
    ([("index", "success"), ("remove", "failed")], False),
    ([("index", "success"), ("remove", "running")], False),
    ([("index", "success"), ("remove", "success"), ("index", "success")], False),
    ([("index", "success"), ("remove", "success"), ("remove", "failed")], True),
    ([("remove", "success"), ("remove", "running")], True),
    ([("remove", "success"), ("index", "failed")], True),
    ([("remove", "success"), ("index", "running")], True),
])
def test_delete_uses_latest_success_and_rechecks_after_detail(settings_factory, operations, allowed):
    client, source, runs = _client(settings_factory)
    summary = _github(client).json()
    candidate = source.ingestion_service.repository.get(summary["candidate_id"])
    path = f"/admin/ingestion/candidates/{candidate.candidate_id}"
    for operation, status in operations:
        # 每次 operation 前读取详情，最后一次 INDEX 成功后 DELETE 必须重新判断。
        assert client.get(path, headers=_headers()).status_code == 200
        run = runs.create_running(
            candidate.candidate_id, source_id=candidate.document.source_id,
            operation=IndexRunOperation(operation),
        )
        if status == "success":
            runs.succeed(run.id)
        elif status == "failed":
            runs.fail(run.id, error_code="FAILED")
    history = runs.list()
    response = client.delete(path, headers=_headers())
    assert response.status_code == (204 if allowed else 409)
    assert runs.list() == history
    assert source.indexing_service.index_calls == source.indexing_service.remove_calls == 0
    if not allowed:
        assert response.json()["error"]["code"] == "CANDIDATE_SOURCE_NOT_REMOVED"
        assert source.ingestion_service.repository.get(candidate.candidate_id) == candidate


def test_delete_all_versions_then_reproposal_and_reindex(settings_factory):
    client, source, runs = _client(settings_factory)
    versions = []
    for content in ("v1", "v2", "v3"):
        candidate = source.ingestion_service.propose(
            document=build_knowledge_document(
                source_type="firecrawl", source_uri="https://example.com/versions",
                title="versions", content=content,
            ), requester=REQUESTER,
        )
        versions.append(candidate)
    for candidate in versions[:2]:
        path = f"/admin/ingestion/candidates/{candidate.candidate_id}"
        assert client.post(f"{path}/approve", headers=_headers(), json={}).status_code == 200
        assert client.post(f"{path}/index", headers=_headers()).status_code == 200
    old_path = f"/admin/ingestion/candidates/{versions[0].candidate_id}"
    assert client.post(f"{old_path}/remove", headers=_headers()).status_code == 200
    history = runs.list()
    with runs._sessions() as session:
        for run in history:
            assert session.get(_IngestionIndexRunRecord, run.id).source_id == versions[0].document.source_id
    for candidate in versions:
        path = f"/admin/ingestion/candidates/{candidate.candidate_id}"
        assert client.get(path, headers=_headers()).json()["can_delete"] is True
        assert client.delete(path, headers=_headers()).status_code == 204
    assert runs.list() == history
    assert source.ingestion_service.repository.load_latest_successfully_indexed() == []
    assert client.get("/admin/ingestion/index-runs", headers=_headers()).json() == [
        item.model_dump(mode="json") for item in history
    ]
    recreated = source.ingestion_service.propose(document=versions[0].document, requester=REQUESTER)
    assert recreated.candidate_id == versions[0].candidate_id
    assert recreated.status is IngestionCandidateStatus.PENDING
    assert source.ingestion_service.repository.load_latest_successfully_indexed() == []
    assert client.get(old_path, headers=_headers()).json()["can_delete"] is True
    assert client.post(f"{old_path}/approve", headers=_headers(), json={}).status_code == 200
    assert client.post(f"{old_path}/index", headers=_headers()).status_code == 200
    assert runs.latest_successful_for_source(recreated.document.source_id).operation is IndexRunOperation.INDEX
    assert client.get(old_path, headers=_headers()).json()["can_delete"] is False
    assert client.delete(old_path, headers=_headers()).status_code == 409
    assert client.post(f"{old_path}/remove", headers=_headers()).status_code == 200
    assert client.delete(old_path, headers=_headers()).status_code == 204


@pytest.mark.parametrize("error,status,code", [
    (RuntimeError("sensitive detail"), 500, "REMOVE_OPERATION_FAILED"),
    (IndexingError("SOURCE_EMBEDDINGS_UNAVAILABLE", "sensitive detail"), 503, "SOURCE_EMBEDDINGS_UNAVAILABLE"),
])
def test_remove_failure_persists_safe_operation(settings_factory, error, status, code):
    client, source, runs = _client(settings_factory)
    candidate = _github(client).json()
    source.indexing_service.remove_error = error
    response = client.post(f"/admin/ingestion/candidates/{candidate['candidate_id']}/remove", headers=_headers())
    assert response.status_code == status
    assert response.json()["error"]["code"] == code
    assert "sensitive" not in response.text
    run = runs.latest(candidate["candidate_id"])
    assert run.operation is IndexRunOperation.REMOVE
    assert run.status is IndexRunStatus.FAILED
    assert run.error_code == run.error_message == code
    assert source.indexing_service.remove_calls == 1


@pytest.mark.parametrize("source_type,uri", [("local_file", "/tmp/retained.txt"), ("web", "https://example.com/retained")])
def test_remove_unsupported_source_fails_closed(settings_factory, monkeypatch, source_type, uri):
    client, deps = _upload_client(settings_factory)
    candidate = deps.ingestion_service.propose(
        document=build_knowledge_document(source_type=source_type, source_uri=uri, title="retained", content="retained"),
        requester=REQUESTER,
    )
    def forbidden(*args, **kwargs):
        pytest.fail("unsupported source must not reach storage deletion")
    monkeypatch.setattr("app.indexing.service.create_record_manager", forbidden)
    monkeypatch.setattr("app.indexing.service.create_vector_store", forbidden)
    response = client.post(f"/admin/ingestion/candidates/{candidate.candidate_id}/remove", headers=_headers())
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "SOURCE_TYPE_UNSUPPORTED"
    assert deps.index_runs.latest(candidate.candidate_id).status is IndexRunStatus.FAILED
    assert deps.ingestion_service.repository.get(candidate.candidate_id) == candidate


def _upload_client(settings_factory, **settings_overrides):
    settings = settings_factory(
        ingestion_admin_api_enabled=True, ingestion_admin_api_key=ADMIN_KEY,
        **settings_overrides,
    )
    deps = get_default_ingestion_dependencies(settings)
    return TestClient(create_app(settings=settings, ingestion_dependencies=deps)), deps


def _upload_bytes(file_type):
    keyword = "UPLOAD_V1_UNIQUE_987654"
    buffer = BytesIO()
    if file_type == "pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        for text in (keyword, "SECOND_PAGE_TEXT"):
            page = writer.add_blank_page(width=300, height=300)
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
            })
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(buffer)
    elif file_type == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph(keyword)
        document.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_CELL_TEXT"
        document.save(buffer)
    else:
        buffer.write(("# Upload\n\n" + keyword).encode())
    return buffer.getvalue()


@pytest.mark.parametrize("file_type", ["txt", "md", "pdf", "docx"])
def test_upload_formats_persist_pending_and_dedupe(settings_factory, file_type):
    client, deps = _upload_client(settings_factory)
    content = _upload_bytes(file_type)
    filename = f"guide.{file_type}"
    first = client.post(
        "/admin/ingestion/sources/upload", headers=_headers(),
        files={"file": (filename, content)}, data={"reason": "review upload"},
    )
    assert first.status_code == 201, first.text
    summary = first.json()
    assert (summary["source_type"], summary["status"], summary["trust_level"]) == (
        "upload", "pending", "untrusted_external",
    )
    assert summary["policy_decision"] == "require_approval"
    assert "content" not in summary
    document = deps.ingestion_service.repository.get(summary["candidate_id"]).document
    assert "UPLOAD_V1_UNIQUE_987654" in document.content
    assert document.metadata["file_name"] == filename
    assert document.metadata["file_type"] == file_type
    assert document.metadata["file_size"] == len(content)
    assert document.source_uri == f"upload://local/{document.metadata['file_sha256']}"
    assert "/tmp/" not in document.model_dump_json()
    if file_type == "pdf":
        assert document.metadata["page_count"] == 2
        assert "[第 2 页]\n\nSECOND_PAGE_TEXT" in document.content
    if file_type == "docx":
        assert "TABLE_CELL_TEXT" in document.content
    duplicate = client.post(
        "/admin/ingestion/sources/upload", headers=_headers(),
        files={"file": (f"renamed.{file_type}", content)},
    )
    assert duplicate.json() == summary
    assert deps.ingestion_service.repository.count() == 1
    assert deps.index_runs.list() == []


@pytest.mark.parametrize("filename,body,status,code", [
    ("bad.exe", b"data", 415, "UPLOAD_UNSUPPORTED_TYPE"),
    ("empty.txt", b"", 422, "UPLOAD_EMPTY_FILE"),
    ("blank.md", b" \n\t", 422, "UPLOAD_NO_TEXT"),
    ("../bad.txt", b"data", 422, "UPLOAD_INVALID_FILENAME"),
    ("bad.pdf", b"fake pdf", 422, "UPLOAD_PARSE_FAILED"),
    ("bad.docx", b"fake docx", 422, "UPLOAD_PARSE_FAILED"),
    ("binary.txt", b"\x00\x01\x02", 422, "UPLOAD_PARSE_FAILED"),
    ("large.txt", b"x" * 33, 413, "UPLOAD_TOO_LARGE"),
])
def test_invalid_upload_returns_safe_4xx_without_candidate(
    settings_factory, filename, body, status, code,
):
    client, deps = _upload_client(settings_factory, ingestion_upload_max_bytes=32)
    response = client.post(
        "/admin/ingestion/sources/upload", headers=_headers(),
        files={"file": (filename, body)},
    )
    assert response.status_code == status, response.text
    assert response.json()["error"]["code"] == code
    assert "/tmp/" not in response.text
    assert deps.ingestion_service.repository.count() == 0


def test_multipart_strips_windows_client_path_before_upload_adapter(settings_factory):
    client, deps = _upload_client(settings_factory)
    response = client.post(
        "/admin/ingestion/sources/upload", headers=_headers(),
        files={"file": ("C:\\Users\\private\\guide.txt", b"safe content")},
    )
    assert response.status_code == 201
    document = deps.ingestion_service.repository.get(response.json()["candidate_id"]).document
    assert document.metadata["file_name"] == document.title == "guide.txt"
    assert "Users" not in document.model_dump_json()


def test_upload_auth_and_multipart_validation(settings_factory):
    client, deps = _upload_client(settings_factory)
    for headers in ({}, _headers("wrong-key")):
        assert client.post(
            "/admin/ingestion/sources/upload", headers=headers,
            files={"file": ("guide.txt", b"body")},
        ).status_code == 401
    assert client.post(
        "/admin/ingestion/sources/upload", headers=_headers(),
        data={"reason": "missing file"},
    ).status_code == 422
    assert client.post(
        "/admin/ingestion/sources/upload", headers=_headers(),
        files={"file": ("guide.txt", b"body")}, data={"reason": "x" * 1001},
    ).status_code == 422
    assert deps.ingestion_service.repository.count() == 0
