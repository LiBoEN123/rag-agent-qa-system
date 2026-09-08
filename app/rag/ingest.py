"""Local Knowledge 文件读取与文本清洗。"""

import re
from pathlib import Path

from charset_normalizer import from_bytes


SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


def read_txt_auto(file_path: str) -> str:
    """自动识别文本编码并读取文件内容。"""
    raw = Path(file_path).read_bytes()
    result = from_bytes(raw).best()
    if result is None:
        raise ValueError(f"无法识别编码: {file_path}")
    return str(result)


def read_pdf_text(file_path: str) -> str:
    """从 PDF 中提取文本内容。"""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError(
            "读取 PDF 需要安装 pypdf，请先运行 `pip install -r requirements.txt`。"
        ) from exc

    reader = PdfReader(file_path)
    pages = []

    for page_idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = clean_text(text)
        if text:
            pages.append(f"[第 {page_idx} 页]\n{text}")

    content = "\n\n".join(pages).strip()
    if not content:
        raise ValueError(f"PDF 未提取到可索引文本，可能是扫描件或图片型 PDF: {file_path}")
    return content


def read_docx_text(file_path: str) -> str:
    """从 DOCX 段落和表格中提取文本内容。"""
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "读取 DOCX 需要安装 python-docx，请先运行 `pip install -r requirements.txt`。"
        ) from exc

    document = Document(file_path)
    parts = []

    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                parts.append(" | ".join(cells))

    content = "\n\n".join(parts).strip()
    if not content:
        raise ValueError(f"DOCX 未提取到可索引文本: {file_path}")
    return content


def read_file_text(file_path: str) -> str:
    """按文件扩展名选择对应的文档解析器。"""
    suffix = Path(file_path).suffix.lower()

    if suffix in {".txt", ".md"}:
        return read_txt_auto(file_path)
    if suffix == ".pdf":
        return read_pdf_text(file_path)
    if suffix == ".docx":
        return read_docx_text(file_path)

    raise ValueError(f"不支持的文件类型: {suffix}")


def collect_supported_files(data_dir: str):
    """收集 data 目录中支持的知识库文件。"""
    root = Path(data_dir)

    if not root.exists():
        return []

    if root.is_file():
        return [root] if root.suffix.lower() in SUPPORTED_SUFFIXES else []

    files = sorted(
        p for p in root.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )
    return files


def clean_text(text: str) -> str:
    """统一换行和空白，得到更稳定的纯文本。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_markdown_text(text: str) -> str:
    """在通用清洗基础上移除 Markdown 注释。"""
    text = clean_text(text)
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    return text.strip()
